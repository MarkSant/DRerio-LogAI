"""Tests for ``BlockDetailDialog.start_session``.

Focused on the "deferred for zone confirmation" vs "genuine failure"
distinction so the user no longer sees a spurious
"Falha ao iniciar sessão para Animal X" popup after approving the
auto-detect polygon — that path returns False from
``start_live_project_session`` because the session is waiting for zone
confirmation, not because anything failed.
"""

from __future__ import annotations

from types import SimpleNamespace
from typing import Any, cast
from unittest.mock import MagicMock, patch

import pandas as pd
from docx import Document

from zebtrack.core.services.session_duration_resolver import duration_override_key
from zebtrack.ui.dialogs.block_detail_dialog import BlockDetailDialog


def _build_dialog(
    *,
    start_session_return: bool,
    pending_zone_confirmation: bool,
    pending_external_trigger: bool = False,
) -> tuple[BlockDetailDialog, MagicMock]:
    """Construct a BlockDetailDialog bypassing Toplevel init.

    Returns (dialog, session_coordinator_mock) so callers can assert on
    the session_coordinator's call history.
    """
    dialog = cast(Any, BlockDetailDialog.__new__(BlockDetailDialog))
    dialog.day_num = 1
    dialog.group_name = "Controle"
    dialog._camera_index_override = None
    dialog._camera_friendly_name_override = None
    dialog.subjects_per_group = 3
    dialog.completed_sessions = set()
    dialog.destroy = MagicMock()  # Skip real Tk teardown
    # ``start_session`` resolve a duração da cobaia antes de delegar.
    dialog.project_manager = SimpleNamespace(
        project_data={"recording_duration_s": 300.0}, project_path=None
    )

    session_coordinator = MagicMock()
    session_coordinator.start_live_project_session.return_value = start_session_return
    session_coordinator.live_calibration_coordinator = SimpleNamespace(
        pending_zone_confirmation=pending_zone_confirmation,
    )
    # Precisa de retorno EXPLÍCITO: um MagicMock cru devolve um filho truthy, o
    # que faria todo start_session parecer "adiado" e engoliria erros reais.
    session_coordinator.has_pending_external_trigger.return_value = pending_external_trigger
    dialog.session_coordinator = session_coordinator

    return cast(BlockDetailDialog, dialog), session_coordinator


def _build_partial_report_dialog(tmp_path) -> BlockDetailDialog:
    dialog = cast(Any, BlockDetailDialog.__new__(BlockDetailDialog))
    dialog.day_num = 1
    dialog.group_name = "Controle"
    dialog.subjects_per_group = 1
    dialog.completed_sessions = {(1, "Controle", "1")}
    dialog.project_manager = SimpleNamespace(
        project_path=str(tmp_path), project_data={"recording_duration_s": 300.0}
    )
    dialog._find_session_folder = MagicMock(return_value=tmp_path / "session_1")
    return cast(BlockDetailDialog, dialog)


def test_start_session_does_not_show_error_when_deferred_for_zones():
    """When ``start_live_project_session`` returns False because zones are
    pending confirmation, the user is on the zone tab seeing the
    "Iniciar Gravação" banner — no error popup should appear."""
    dialog, _ = _build_dialog(start_session_return=False, pending_zone_confirmation=True)

    with patch("zebtrack.ui.dialogs.block_detail_dialog.messagebox.showerror") as mock_error:
        dialog.start_session("1")

    mock_error.assert_not_called()


def test_start_session_shows_error_on_genuine_failure():
    """A return of False with ``pending_zone_confirmation=False`` is a real
    failure (camera missing, wrong project type, etc.) — the popup must
    still surface so the user knows the session didn't start."""
    dialog, _ = _build_dialog(start_session_return=False, pending_zone_confirmation=False)

    with patch("zebtrack.ui.dialogs.block_detail_dialog.messagebox.showerror") as mock_error:
        dialog.start_session("1")

    mock_error.assert_called_once()
    args, _kwargs = mock_error.call_args
    # Body must reference the subject + day/group so the user can act on it.
    assert "Animal 1" in args[1]
    assert "Day 1" in args[1]
    assert "Controle" in args[1]


def test_start_session_shows_no_error_on_success():
    """Happy path: success=True → no error popup regardless of pending flag."""
    dialog, _ = _build_dialog(start_session_return=True, pending_zone_confirmation=False)

    with patch("zebtrack.ui.dialogs.block_detail_dialog.messagebox.showerror") as mock_error:
        dialog.start_session("1")

    mock_error.assert_not_called()


def test_start_session_handles_missing_live_calibration_coordinator():
    """Older session_coordinator builds may not expose
    ``live_calibration_coordinator``. The deferred-detection probe must
    tolerate that and fall back to showing the error popup."""
    dialog, session_coordinator = _build_dialog(
        start_session_return=False, pending_zone_confirmation=False
    )
    # Strip the attribute to simulate a legacy build.
    del session_coordinator.live_calibration_coordinator

    with patch("zebtrack.ui.dialogs.block_detail_dialog.messagebox.showerror") as mock_error:
        dialog.start_session("1")

    mock_error.assert_called_once()


def test_get_session_files_status_detects_relatorio_excel(tmp_path):
    dialog = _build_partial_report_dialog(tmp_path)
    session_folder = tmp_path / "session_1"
    session_folder.mkdir()
    (session_folder / "4_Relatorio_live_exp.xlsx").touch()

    status = dialog._get_session_files_status(session_folder)

    assert status["summary"] is True


def test_generate_partial_report_accepts_relatorio_excel(tmp_path):
    dialog = _build_partial_report_dialog(tmp_path)
    session_folder = tmp_path / "session_1"
    session_folder.mkdir()
    pd.DataFrame({"total_distance_cm": [10.0]}).to_excel(
        session_folder / "4_Relatorio_live_exp.xlsx",
        index=False,
    )
    dialog.session_coordinator = cast(Any, SimpleNamespace(event_bus=None))
    dialog.live_batch_coordinator = cast(Any, SimpleNamespace(event_bus=None))

    with (
        patch("zebtrack.ui.dialogs.block_detail_dialog.messagebox.showwarning") as mock_warning,
        patch("zebtrack.ui.dialogs.block_detail_dialog.messagebox.showinfo") as mock_info,
        patch("zebtrack.ui.dialogs.block_detail_dialog.messagebox.showerror") as mock_error,
        patch(
            "zebtrack.ui.dialogs.block_detail_dialog.messagebox.askyesno",
            side_effect=[False, False],
        ),
    ):
        dialog.generate_partial_report()

    mock_warning.assert_not_called()
    mock_error.assert_not_called()
    mock_info.assert_called_once()
    assert (tmp_path / "partial_reports" / "PartialReport_Dia1_Controle.xlsx").exists()
    assert (tmp_path / "partial_reports" / "PartialReport_Dia1_Controle.docx").exists()


def test_get_partial_report_stats_columns_ignores_non_numeric_time_fields(tmp_path):
    dialog = _build_partial_report_dialog(tmp_path)
    df = pd.DataFrame(
        {
            "analysis_timestamp": ["2026-06-07 10:00:00"],
            "total_time_seconds": [12.5],
            "total_roi_entries": [3],
            "time_label": ["baseline"],
            "animal": ["1"],
        }
    )

    stats_cols = dialog._get_partial_report_stats_columns(df)

    assert stats_cols == ["total_time_seconds", "total_roi_entries"]


def test_write_partial_report_word_lists_only_successfully_parsed_sessions(tmp_path):
    dialog = _build_partial_report_dialog(tmp_path)
    session_folder = tmp_path / "session_1"
    session_folder.mkdir()

    valid_summary = session_folder / "4_Relatorio_live_exp.xlsx"
    invalid_summary = session_folder / "4_Relatorio_corrompido.xlsx"

    pd.DataFrame({"total_distance_cm": [10.0]}).to_excel(valid_summary, index=False)
    invalid_summary.write_text("not an excel file", encoding="utf-8")

    all_data, unified_df, parsed_summary_files = dialog._build_partial_report_dataset(
        [("1", valid_summary), ("2", invalid_summary)]
    )

    output_path = tmp_path / "partial.docx"
    dialog._write_partial_report_word(
        output_path,
        "partial.xlsx",
        parsed_summary_files,
        all_data,
        unified_df,
    )

    document = Document(output_path)
    session_rows = document.tables[0].rows

    assert len(session_rows) == 2
    assert session_rows[1].cells[0].text == "1"
    assert session_rows[1].cells[1].text == valid_summary.name


# ---------------------------------------------------------------------------
# Duração por bloco/cobaia + gatilho externo
# ---------------------------------------------------------------------------


def test_start_session_passes_resolved_duration():
    """A duração resolvida tem de CHEGAR ao coordinator — sem isso o override
    aparece na UI e a gravação continua usando o padrão do projeto."""
    dialog, session_coordinator = _build_dialog(
        start_session_return=True, pending_zone_confirmation=False
    )
    dialog.project_manager.project_data = {
        "recording_duration_s": 300.0,
        "session_duration_overrides": {duration_override_key(1, "Controle", "1"): 900.0},
    }

    dialog.start_session("1")

    _args, kwargs = session_coordinator.start_live_project_session.call_args
    assert kwargs["duration_s"] == 900.0


def test_start_session_falls_back_to_project_duration():
    dialog, session_coordinator = _build_dialog(
        start_session_return=True, pending_zone_confirmation=False
    )
    dialog.project_manager.project_data = {"recording_duration_s": 420.0}

    dialog.start_session("2")

    _args, kwargs = session_coordinator.start_live_project_session.call_args
    assert kwargs["duration_s"] == 420.0


def test_start_session_no_error_when_armed_for_external_trigger():
    """Sessão armada aguardando o Arduino retorna False, mas NÃO é falha:
    o usuário já vê o aviso 'Aguardando sinal externo'."""
    dialog, _ = _build_dialog(
        start_session_return=False,
        pending_zone_confirmation=False,
        pending_external_trigger=True,
    )

    with patch("zebtrack.ui.dialogs.block_detail_dialog.messagebox.showerror") as mock_error:
        dialog.start_session("1")

    mock_error.assert_not_called()


def test_heterogeneous_duration_warning_absent_when_uniform(tmp_path):
    dialog = _build_partial_report_dialog(tmp_path)
    assert dialog._heterogeneous_duration_warning(["1", "2"]) is None


def test_heterogeneous_duration_warning_lists_distinct_durations(tmp_path):
    dialog = _build_partial_report_dialog(tmp_path)
    dialog.project_manager.project_data = {
        "recording_duration_s": 300.0,
        "session_duration_overrides": {duration_override_key(1, "Controle", "2"): 900.0},
    }

    warning = dialog._heterogeneous_duration_warning(["1", "2"])

    assert warning is not None
    assert "5 min" in warning
    assert "15 min" in warning
    assert "video_duration_s" in warning


def test_partial_report_stats_columns_include_duration():
    """``video_duration_s`` não casa com nenhum dos keywords antigos; sem ela o
    agregado esconde o dado que torna as métricas comparáveis (ou não)."""
    df = pd.DataFrame(
        {
            "video_duration_s": [300.0, 900.0],
            "total_distance": [1.0, 2.0],
            "analysis_timestamp": [1, 2],
        }
    )

    columns = BlockDetailDialog._get_partial_report_stats_columns(df)

    assert "video_duration_s" in columns
    assert "analysis_timestamp" not in columns


# ---------------------------------------------------------------------------
# Notas experimentais (bloco e cobaia)
# ---------------------------------------------------------------------------


def _build_note_dialog(tmp_path, project_data: dict | None = None):
    """Dialog wired just enough to drive the note flow headless.

    Returns ``(dialog, calls)`` where ``calls`` records save/popup ORDER, which
    is the contract under test: success may only be announced after the save.
    """
    calls = MagicMock()
    dialog = cast(Any, BlockDetailDialog.__new__(BlockDetailDialog))
    dialog.day_num = 1
    dialog.group_name = "Controle"
    dialog.subjects_per_group = 3
    dialog.completed_sessions = set()
    dialog.project_manager = SimpleNamespace(
        project_path=str(tmp_path),
        project_data=project_data if project_data is not None else {},
        save_project=calls.save_project,
    )
    dialog.session_coordinator = SimpleNamespace(event_bus=calls.event_bus)
    dialog.live_batch_coordinator = SimpleNamespace(event_bus=None)
    dialog._block_note_label = MagicMock()
    dialog._note_button = MagicMock()
    dialog.refresh_subject_rows = calls.refresh_subject_rows
    return cast(BlockDetailDialog, dialog), calls


_MB = "zebtrack.ui.dialogs.block_detail_dialog.messagebox"


def test_add_note_saves_before_announcing_success(tmp_path):
    dialog, calls = _build_note_dialog(tmp_path)
    dialog._ask_note_text = MagicMock(return_value="  luz piscou\nno minuto 3 ")

    with (
        patch(f"{_MB}.showinfo", calls.showinfo),
        patch(f"{_MB}.showerror", calls.showerror),
    ):
        dialog.add_note()

    names = [c[0] for c in calls.mock_calls]
    assert names.index("save_project") < names.index("showinfo")
    calls.showerror.assert_not_called()
    stored = dialog.project_manager.project_data["experiment_notes"]["Dia_1|Controle|*"]
    assert stored["text"] == "luz piscou\nno minuto 3"
    dialog._block_note_label.config.assert_called_once()
    dialog._note_button.config.assert_called_once_with(text="📝 Edit Note")


def test_add_note_save_failure_reports_error_and_rolls_back(tmp_path):
    previous = {"Dia_1_Controle": "nota antiga"}
    dialog, calls = _build_note_dialog(tmp_path, {"experiment_notes": dict(previous)})
    calls.save_project.side_effect = OSError("disk full")
    dialog._ask_note_text = MagicMock(return_value="nova")

    with (
        patch(f"{_MB}.showinfo", calls.showinfo),
        patch(f"{_MB}.showerror", calls.showerror),
    ):
        dialog.add_note()

    calls.showinfo.assert_not_called()
    calls.showerror.assert_called_once()
    assert "note could not be written" in calls.showerror.call_args[0][1]
    # Nothing half-written stays in memory for an unrelated save to persist later.
    assert dialog.project_manager.project_data["experiment_notes"] == previous
    calls.event_bus.publish.assert_not_called()


def test_add_note_rollback_removes_container_that_did_not_exist(tmp_path):
    dialog, calls = _build_note_dialog(tmp_path)
    calls.save_project.side_effect = OSError("disk full")
    dialog._ask_note_text = MagicMock(return_value="nova")

    with patch(f"{_MB}.showerror"):
        dialog.add_note()

    assert "experiment_notes" not in dialog.project_manager.project_data


def test_add_note_without_project_path_does_not_save(tmp_path):
    dialog, calls = _build_note_dialog(tmp_path)
    dialog.project_manager.project_path = None
    dialog._ask_note_text = MagicMock(return_value="nova")

    with (
        patch(f"{_MB}.showinfo", calls.showinfo),
        patch(f"{_MB}.showerror", calls.showerror),
    ):
        dialog.add_note()

    calls.save_project.assert_not_called()
    calls.showinfo.assert_not_called()
    assert "no path defined, so the note" in calls.showerror.call_args[0][1]


def test_add_note_cancel_changes_nothing(tmp_path):
    dialog, calls = _build_note_dialog(tmp_path, {"experiment_notes": {"Dia_1_Controle": "x"}})
    dialog._ask_note_text = MagicMock(return_value=None)

    dialog.add_note()

    calls.save_project.assert_not_called()
    assert dialog.project_manager.project_data == {"experiment_notes": {"Dia_1_Controle": "x"}}


def test_add_note_prefills_legacy_note_and_blank_removes_it(tmp_path):
    dialog, calls = _build_note_dialog(tmp_path, {"experiment_notes": {"Dia_1_Controle": "antiga"}})
    dialog._ask_note_text = MagicMock(return_value="")

    with patch(f"{_MB}.showinfo", calls.showinfo):
        dialog.add_note()

    assert dialog._ask_note_text.call_args[0][2] == "antiga"
    assert dialog.project_manager.project_data["experiment_notes"] == {}
    assert calls.showinfo.call_args[0][0] == "Note Removed"
    dialog._note_button.config.assert_called_once_with(text="📝 Add Note")


def test_subject_note_is_saved_and_rows_redrawn(tmp_path):
    dialog, calls = _build_note_dialog(tmp_path)
    dialog._ask_note_text = MagicMock(return_value="pulou do aquario")

    dialog._edit_subject_note("2")

    stored = dialog.project_manager.project_data["experiment_notes"]["Dia_1|Controle|2"]
    assert stored["text"] == "pulou do aquario"
    calls.save_project.assert_called_once()
    calls.refresh_subject_rows.assert_called_once()
    calls.event_bus.publish.assert_called_once()


def test_subject_note_not_redrawn_when_save_fails(tmp_path):
    dialog, calls = _build_note_dialog(tmp_path)
    calls.save_project.side_effect = OSError("disk full")
    dialog._ask_note_text = MagicMock(return_value="pulou")

    with patch(f"{_MB}.showerror"):
        dialog._edit_subject_note("2")

    calls.refresh_subject_rows.assert_not_called()
    assert "experiment_notes" not in dialog.project_manager.project_data


def test_format_block_note_shows_preview_and_date(tmp_path):
    dialog, _calls = _build_note_dialog(
        tmp_path,
        {
            "experiment_notes": {
                "Dia_1|Controle|*": {
                    "text": "primeira linha\nsegunda",
                    "updated_at": "2026-10-22T14:30:05",
                }
            }
        },
    )

    assert dialog._format_block_note() == "primeira linha …  (22/10/2026 14:30)"
    assert dialog._note_button_text() == "📝 Edit Note"


def test_format_block_note_without_note(tmp_path):
    dialog, _calls = _build_note_dialog(tmp_path)

    assert dialog._format_block_note() == "(none)"
    assert dialog._note_button_text() == "📝 Add Note"


def test_note_preview_truncates_long_first_line():
    from zebtrack.core.services.experiment_notes import ExperimentNote
    from zebtrack.ui.dialogs.block_detail_dialog import _note_preview

    assert _note_preview(ExperimentNote("curta")) == "curta"
    assert _note_preview(ExperimentNote("x" * 200), limit=10) == "x" * 9 + " …"


def _notes_project_data() -> dict:
    return {
        "recording_duration_s": 300.0,
        "experiment_notes": {
            "Dia_1|Controle|*": {"text": "luz piscou", "updated_at": "2026-10-22T14:30:05"},
            "Dia_1|Controle|1": {"text": "nadou pouco", "updated_at": "2026-10-22T15:00:00"},
            "Dia_1|Controle|3": {"text": "camera caiu, sessao perdida"},
            "Dia_1|CBD|1": {"text": "outro grupo"},
        },
    }


def _write_partial_word(tmp_path, dialog) -> Any:
    session_folder = tmp_path / "session_1"
    session_folder.mkdir()
    summary = session_folder / "4_Relatorio_live_exp.xlsx"
    pd.DataFrame({"total_distance_cm": [10.0]}).to_excel(summary, index=False)
    all_data, unified_df, parsed = dialog._build_partial_report_dataset([("1", summary)])
    output_path = tmp_path / "partial.docx"
    dialog._write_partial_report_word(output_path, "partial.xlsx", parsed, all_data, unified_df)
    return Document(output_path)


def test_partial_word_report_includes_block_and_subject_notes(tmp_path):
    dialog = _build_partial_report_dialog(tmp_path)
    dialog.project_manager.project_data = _notes_project_data()

    document = _write_partial_word(tmp_path, dialog)

    paragraphs = [p.text for p in document.paragraphs]
    assert any("see 'Experimental notes' at the end" in p for p in paragraphs)
    assert "Experimental notes" in paragraphs
    assert "Block note (last edited 22/10/2026 14:30):" in paragraphs
    assert "luz piscou" in paragraphs

    rows = [[cell.text for cell in row.cells] for row in document.tables[-1].rows]
    assert rows == [
        ["Animal", "Note", "Last edited"],
        ["1", "nadou pouco", "22/10/2026 15:00"],
        ["3 (not in this report)", "camera caiu, sessao perdida", "-"],
    ]
    # The session table stays the first table, untouched by the notes section.
    assert document.tables[0].rows[0].cells[1].text == "Source file"


def test_partial_word_report_without_notes_has_no_notes_section(tmp_path):
    dialog = _build_partial_report_dialog(tmp_path)

    document = _write_partial_word(tmp_path, dialog)

    paragraphs = [p.text for p in document.paragraphs]
    assert "Experimental notes" not in paragraphs
    assert not any("see 'Experimental notes'" in p for p in paragraphs)


def test_partial_excel_report_has_notes_sheet_only_with_notes(tmp_path):
    dialog = _build_partial_report_dialog(tmp_path)
    unified_df = pd.DataFrame({"animal": ["1"], "total_distance_cm": [10.0]})

    without_notes = tmp_path / "sem_notas.xlsx"
    dialog._write_partial_report_excel(without_notes, [unified_df], unified_df)
    with pd.ExcelFile(without_notes) as workbook:
        assert "Notas" not in workbook.sheet_names

    dialog.project_manager.project_data = _notes_project_data()
    with_notes = tmp_path / "com_notas.xlsx"
    dialog._write_partial_report_excel(with_notes, [unified_df], unified_df)

    notes = pd.read_excel(with_notes, sheet_name="Notas", dtype=str, keep_default_na=False)
    assert notes.to_dict("records") == [
        {
            "day": "Dia_1",
            "group": "Controle",
            "scope": "block",
            "subject": "",
            "note": "luz piscou",
            "updated_at": "2026-10-22T14:30:05",
        },
        {
            "day": "Dia_1",
            "group": "Controle",
            "scope": "subject",
            "subject": "1",
            "note": "nadou pouco",
            "updated_at": "2026-10-22T15:00:00",
        },
        {
            "day": "Dia_1",
            "group": "Controle",
            "scope": "subject",
            "subject": "3",
            "note": "camera caiu, sessao perdida",
            "updated_at": "",
        },
    ]
